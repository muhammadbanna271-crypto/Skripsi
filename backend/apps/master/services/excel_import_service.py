"""Import master data dari file Excel (Kecamatan; Desa; Variabel + Indikator +
Pernyataan).

Format file mengikuti folder ``Data`` (Kecamatan_TRIP.xlsx, Desa_TRIP.xlsx,
Data_Variabel.xlsx).
Kolom dibaca berdasarkan URUTAN (A, B, C, ...), bukan nama header, supaya
file asli bisa langsung dipakai. Baris pertama dianggap header kalau
isinya bukan data.

Import bersifat all-or-nothing: kalau ada satu baris bermasalah, tidak ada
data yang disimpan dan semua masalah dilaporkan sekaligus per baris.
"""

import re
from collections import OrderedDict
from io import BytesIO

from django.db import transaction
from django.db.models import Max
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from apps.master.models import (
    District,
    Indicator,
    Questionnaire,
    Variable,
    Village,
)
from apps.master.services.variable_configuration_service import (
    VariableConfigurationService,
)

MAX_ROWS = 5000
MAX_ERRORS_SHOWN = 100

VARIABLE_CODE = re.compile(r"^([XYZ])(\d+)$")
INDICATOR_CODE = re.compile(r"^([XYZ])(\d+)\.(\d+)$")

ROLE_BY_PREFIX = {
    "X": Variable.ROLE_PREDICTOR,
    "Y": Variable.ROLE_MEDIATOR,
    "Z": Variable.ROLE_RESPONSE,
}

# Kolom "Jenis Jawaban" -> Questionnaire.answer_type. Pencocokan tidak
# peka huruf besar/kecil maupun spasi.
ANSWER_TYPE_LABELS = OrderedDict([
    ("likert", "Skala Likert (1-5)"),
    ("boolean", "Ya/Tidak"),
    ("integer", "Angka Bulat"),
    ("decimal", "Desimal"),
    ("text", "Teks"),
    ("choice", "Pilihan Ganda"),
])
_ANSWER_ALIASES = {
    "likert": ["likert", "skalalikert(1-5)", "likertscale(1-5)", "skalalikert"],
    "boolean": ["boolean", "ya/tidak", "yes/no"],
    "integer": ["integer", "angkabulat", "bilanganbulat"],
    "decimal": ["decimal", "desimal", "angkadesimal"],
    "text": ["text", "teks", "isianteks"],
    "choice": ["choice", "pilihanganda", "multiplechoice"],
}
ANSWER_TYPE_LOOKUP = {
    alias: key for key, aliases in _ANSWER_ALIASES.items() for alias in aliases
}


def parse_answer_type(label):
    return ANSWER_TYPE_LOOKUP.get("".join(label.split()).casefold())


# Mediator wajib punya layer; file Excel tidak membawa informasi layer,
# jadi semua mediator hasil import masuk layer 1 (bisa dipindah lewat
# halaman Configure Variables).
DEFAULT_MEDIATOR_LAYER = 1


class ExcelImportError(Exception):
    """File tidak bisa dibaca sama sekali (bukan Excel, kosong, terlalu besar)."""


# ======================================================================
# Pembacaan file
# ======================================================================

def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _read_rows(uploaded_file, columns):
    """Kembalikan list (nomor_baris_excel, [nilai kolom 1..columns])."""
    name = getattr(uploaded_file, "name", "") or ""
    if not name.lower().endswith(".xlsx"):
        raise ExcelImportError("File harus berformat Excel (.xlsx).")

    try:
        workbook = load_workbook(
            filename=BytesIO(uploaded_file.read()),
            read_only=True,
            data_only=True,
        )
    except Exception as error:  # BadZipFile, InvalidFileException, dst.
        raise ExcelImportError(
            "File tidak dapat dibaca sebagai Excel (.xlsx)."
        ) from error

    sheet = workbook.worksheets[0]
    rows = []
    for number, row in enumerate(
        sheet.iter_rows(min_row=1, max_col=columns, values_only=True),
        start=1,
    ):
        values = [_clean(v) for v in row]
        values += [""] * (columns - len(values))
        if not any(values):
            continue
        rows.append((number, values))
        if len(rows) > MAX_ROWS:
            raise ExcelImportError(
                f"File terlalu besar (maksimal {MAX_ROWS} baris data)."
            )
    workbook.close()

    if not rows:
        raise ExcelImportError("File kosong: tidak ada baris data.")
    return rows


# ======================================================================
# Hasil import
# ======================================================================

class ImportResult:
    def __init__(self):
        self.errors = []     # list (baris, pesan)
        self.created = OrderedDict()
        self.skipped = OrderedDict()

    @property
    def ok(self):
        return not self.errors

    def error(self, row, message):
        self.errors.append((row, message))

    def count(self, bucket, key, amount=1):
        target = self.created if bucket == "created" else self.skipped
        target[key] = target.get(key, 0) + amount

    @property
    def errors_shown(self):
        return self.errors[:MAX_ERRORS_SHOWN]

    @property
    def errors_hidden(self):
        return max(0, len(self.errors) - MAX_ERRORS_SHOWN)


# ======================================================================
# Kecamatan
# ======================================================================

class DistrictExcelImportService:
    """Kolom: A = ID Kecamatan, B = Kecamatan."""

    COLUMNS = 2

    @classmethod
    def import_file(cls, uploaded_file):
        rows = _read_rows(uploaded_file, cls.COLUMNS)

        if not rows[0][1][0].isdigit():
            rows = rows[1:]
        if not rows:
            raise ExcelImportError("File hanya berisi header, tidak ada data kecamatan.")

        result = ImportResult()
        parsed = cls._validate(rows, result)
        if not result.ok:
            return result

        with transaction.atomic():
            for number, code, name in parsed:
                if District.objects.filter(code=code).exists():
                    result.count("skipped", "Kecamatan (sudah ada, tidak berubah)")
                else:
                    District.objects.create(code=code, name=name)
                    result.count("created", "Kecamatan")
        return result

    @staticmethod
    def _validate(rows, result):
        parsed = []
        seen_codes = {}
        seen_names = {}
        code_max = District._meta.get_field("code").max_length
        name_max = District._meta.get_field("name").max_length

        for number, (code, name) in rows:
            problems = len(result.errors)
            if not code.isdigit():
                result.error(number, f"ID Kecamatan '{code}' harus berupa angka.")
            elif len(code) > code_max:
                result.error(number, f"ID Kecamatan '{code}' terlalu panjang.")
            if not name:
                result.error(number, "Nama Kecamatan tidak boleh kosong.")
            elif len(name) > name_max:
                result.error(number, f"Nama Kecamatan maksimal {name_max} karakter.")
            if len(result.errors) > problems:
                continue

            code = str(int(code))
            if code in seen_codes:
                result.error(
                    number,
                    f"ID Kecamatan {code} sudah dipakai di baris {seen_codes[code]}.",
                )
                continue
            if name.casefold() in seen_names:
                result.error(
                    number,
                    f"Kecamatan '{name}' sudah ada di baris "
                    f"{seen_names[name.casefold()]}.",
                )
                continue
            seen_codes[code] = number
            seen_names[name.casefold()] = number
            parsed.append((number, code, name))

        for number, code, name in parsed:
            by_code = District.objects.filter(code=code).first()
            by_name = District.objects.filter(name__iexact=name).first()
            if by_code and by_code.name.casefold() != name.casefold():
                result.error(
                    number,
                    f"ID Kecamatan {code} sudah dipakai '{by_code.name}' di sistem.",
                )
            elif by_name and by_name.code != code:
                result.error(
                    number,
                    f"Kecamatan '{name}' sudah ada di sistem dengan ID {by_name.code}.",
                )
        return parsed


# ======================================================================
# Desa
# ======================================================================

class VillageExcelImportService:
    """Kolom: A = ID Desa, B = Kecamatan, C = Desa/Kelurahan."""

    COLUMNS = 3

    @classmethod
    def import_file(cls, uploaded_file):
        rows = _read_rows(uploaded_file, cls.COLUMNS)

        # Header = baris pertama kalau ID Desa-nya bukan angka.
        if not rows[0][1][0].isdigit():
            rows = rows[1:]
        if not rows:
            raise ExcelImportError("File hanya berisi header, tidak ada data desa.")

        result = ImportResult()
        parsed = cls._validate(rows, result)
        if not result.ok:
            return result

        with transaction.atomic():
            cls._save(parsed, result)
        return result

    @staticmethod
    def _validate(rows, result):
        parsed = []
        seen_codes = {}
        seen_names = {}

        for number, (code, district, name) in rows:
            problems = len(result.errors)

            if not code.isdigit():
                result.error(number, f"ID Desa '{code}' harus berupa angka.")
            elif len(code) > Village._meta.get_field("code").max_length:
                result.error(number, f"ID Desa '{code}' terlalu panjang.")
            if not district:
                result.error(number, "Kecamatan tidak boleh kosong.")
            elif len(district) > District._meta.get_field("name").max_length:
                result.error(number, "Nama kecamatan terlalu panjang.")
            if not name:
                result.error(number, "Nama Desa/Kelurahan tidak boleh kosong.")
            elif len(name) > Village._meta.get_field("name").max_length:
                result.error(number, "Nama desa terlalu panjang.")

            if len(result.errors) > problems:
                continue

            code = str(int(code))
            if code in seen_codes:
                result.error(
                    number,
                    f"ID Desa {code} sudah dipakai di baris {seen_codes[code]}.",
                )
                continue
            if name.casefold() in seen_names:
                result.error(
                    number,
                    f"Nama desa '{name}' sudah ada di baris "
                    f"{seen_names[name.casefold()]}.",
                )
                continue
            seen_codes[code] = number
            seen_names[name.casefold()] = number
            parsed.append((number, code, district, name))

        # Konflik dengan data yang sudah ada di database.
        existing_by_code = {v.code: v for v in Village.objects.all()}
        existing_by_name = {v.name.casefold(): v for v in Village.objects.all()}
        for number, code, district, name in parsed:
            by_code = existing_by_code.get(code)
            by_name = existing_by_name.get(name.casefold())
            if by_code and by_code.name.casefold() != name.casefold():
                result.error(
                    number,
                    f"ID Desa {code} sudah dipakai desa '{by_code.name}' di sistem.",
                )
            elif by_name and by_name.code != code:
                result.error(
                    number,
                    f"Desa '{name}' sudah ada di sistem dengan ID {by_name.code}.",
                )
        return parsed

    @staticmethod
    def _next_district_code():
        numbers = [
            int(code)
            for code in District.objects.values_list("code", flat=True)
            if code.isdigit()
        ]
        return str(max(numbers, default=0) + 1)

    @classmethod
    def _save(cls, parsed, result):
        districts = {d.name.casefold(): d for d in District.objects.all()}

        for number, code, district_name, name in parsed:
            district = districts.get(district_name.casefold())
            if district is None:
                district = District.objects.create(
                    code=cls._next_district_code(), name=district_name
                )
                districts[district_name.casefold()] = district
                result.count("created", "Kecamatan")

            village = Village.objects.filter(code=code).first()
            if village is None:
                Village.objects.create(code=code, name=name, district=district)
                result.count("created", "Desa")
            elif village.district_id != district.id:
                village.district = district
                village.save(update_fields=["district", "updated_at"])
                result.count("created", "Desa (kecamatan diperbarui)")
            else:
                result.count("skipped", "Desa (sudah ada, tidak berubah)")


# ======================================================================
# Variabel + Indikator + Pernyataan
# ======================================================================

class VariableExcelImportService:
    """Kolom: A = Kode Variabel, B = Nama Variabel, C = Kode Indikator,
    D = Nama Indikator, E = Pernyataan (menjadi pertanyaan kuesioner),
    F = Jenis Jawaban (mis. Skala Likert (1-5))."""

    COLUMNS = 6

    @classmethod
    def import_file(cls, uploaded_file):
        rows = _read_rows(uploaded_file, cls.COLUMNS)

        if not VARIABLE_CODE.match(rows[0][1][0].upper()):
            rows = rows[1:]
        if not rows:
            raise ExcelImportError("File hanya berisi header, tidak ada data.")

        result = ImportResult()
        parsed = cls._validate(rows, result)
        if not result.ok:
            return result

        with transaction.atomic():
            cls._save(parsed, result)
        return result

    @staticmethod
    def _validate(rows, result):
        parsed = []
        variable_names = {}   # kode variabel -> nama
        name_owner = {}       # nama variabel (casefold) -> kode
        indicator_codes = {}  # kode indikator -> baris
        indicator_names = {}  # nama indikator (casefold) -> kode indikator

        var_max = Variable._meta.get_field("name").max_length
        ind_max = Indicator._meta.get_field("name").max_length

        for number, (v_code, v_name, i_code, i_name, statement, answer) in rows:
            v_code = v_code.upper()
            i_code = i_code.upper()
            problems = len(result.errors)

            v_match = VARIABLE_CODE.match(v_code)
            i_match = INDICATOR_CODE.match(i_code)

            if not v_match:
                result.error(
                    number,
                    f"Kode Variabel '{v_code}' tidak valid (contoh: X1, Y2, Z1).",
                )
            if not i_match:
                result.error(
                    number,
                    f"Kode Indikator '{i_code}' tidak valid (contoh: X1.1).",
                )
            elif v_match and f"{i_match[1]}{i_match[2]}" != v_code:
                result.error(
                    number,
                    f"Kode Indikator {i_code} tidak sesuai dengan Variabel {v_code}.",
                )
            if not v_name:
                result.error(number, "Nama Variabel tidak boleh kosong.")
            elif len(v_name) > var_max:
                result.error(number, f"Nama Variabel maksimal {var_max} karakter.")
            if not i_name:
                result.error(number, "Nama Indikator tidak boleh kosong.")
            elif len(i_name) > ind_max:
                result.error(number, f"Nama Indikator maksimal {ind_max} karakter.")
            if not statement:
                result.error(number, "Pernyataan tidak boleh kosong.")
            answer_type = parse_answer_type(answer)
            if answer_type is None:
                allowed = ", ".join(ANSWER_TYPE_LABELS.values())
                result.error(
                    number,
                    f"Jenis Jawaban '{answer}' tidak dikenal "
                    f"(pilihan: {allowed}).",
                )

            if len(result.errors) > problems:
                continue

            if variable_names.setdefault(v_code, v_name) != v_name:
                result.error(
                    number,
                    f"Variabel {v_code} punya dua nama berbeda: "
                    f"'{variable_names[v_code]}' dan '{v_name}'.",
                )
                continue
            if name_owner.setdefault(v_name.casefold(), v_code) != v_code:
                result.error(
                    number,
                    f"Nama Variabel '{v_name}' dipakai juga oleh "
                    f"{name_owner[v_name.casefold()]}.",
                )
                continue
            if i_code in indicator_codes:
                result.error(
                    number,
                    f"Kode Indikator {i_code} sudah muncul di baris "
                    f"{indicator_codes[i_code]}.",
                )
                continue
            if indicator_names.setdefault(i_name.casefold(), i_code) != i_code:
                result.error(
                    number,
                    f"Nama Indikator '{i_name}' dipakai juga oleh "
                    f"{indicator_names[i_name.casefold()]}.",
                )
                continue

            indicator_codes[i_code] = number
            parsed.append(
                (number, v_match[1], int(v_match[2]), v_code, v_name,
                 int(i_match[3]), i_code, i_name, statement, answer_type)
            )

        # Konflik dengan data yang sudah ada: nama indikator yang sama
        # tidak boleh menempel pada variabel lain.
        existing_vars = {v.name.casefold(): v for v in Variable.objects.all()}
        for entry in parsed:
            number, _, _, v_code, v_name, _, i_code, i_name, _, _ = entry
            existing = (
                Indicator.objects.select_related("variable")
                .filter(name__iexact=i_name)
                .first()
            )
            if existing is None:
                continue
            owner = existing_vars.get(v_name.casefold())
            if owner is None or existing.variable_id != owner.id:
                result.error(
                    number,
                    f"Indikator '{i_name}' sudah ada di sistem pada variabel "
                    f"'{existing.variable.name}'.",
                )
        return parsed

    @classmethod
    def _save(cls, parsed, result):
        # Urutan variabel baru mengikuti nomor pada kode file (X1, X2, ...),
        # lalu ditambahkan di akhir urutan role-nya.
        parsed = sorted(parsed, key=lambda e: (e[1], e[2], e[5], e[0]))

        variables = {}
        temp_counter = 0
        for entry in parsed:
            (number, prefix, v_num, v_code, v_name, i_num, i_code, i_name,
             statement, answer_type) = entry

            variable = variables.get(v_code)
            if variable is None:
                variable = Variable.objects.filter(name__iexact=v_name).first()
                if variable is None:
                    variable = cls._create_variable(prefix, v_name)
                    result.count("created", "Variabel")
                else:
                    result.count("skipped", "Variabel (sudah ada, tidak berubah)")
                variables[v_code] = variable

            indicator = Indicator.objects.filter(name__iexact=i_name).first()
            if indicator is None:
                # Kode sementara; regenerate_codes() memberi kode final
                # "<kode variabel>.<nomor>" tanpa bentrok dengan data lama.
                temp_counter += 1
                indicator = Indicator.objects.create(
                    variable=variable,
                    name=i_name,
                    code=f"_{temp_counter}.{i_num}",
                )
                result.count("created", "Indikator")
            else:
                result.count("skipped", "Indikator (sudah ada, tidak berubah)")

            if Questionnaire.objects.filter(
                indicator=indicator, question__iexact=statement
            ).exists():
                result.count("skipped", "Pernyataan (sudah ada)")
                continue
            last = Questionnaire.objects.filter(indicator=indicator).aggregate(
                m=Max("question_order")
            )["m"]
            Questionnaire.objects.create(
                indicator=indicator,
                question=statement,
                answer_type=answer_type,
                question_order=(last or 0) + 1,
            )
            result.count("created", "Pernyataan")

        if result.created.get("Variabel") or result.created.get("Indikator"):
            VariableConfigurationService.regenerate_codes()

    @staticmethod
    def _create_variable(prefix, name):
        role = ROLE_BY_PREFIX[prefix]
        layer = None
        if role == Variable.ROLE_MEDIATOR:
            layer = VariableConfigurationService._get_or_create_layer(
                DEFAULT_MEDIATOR_LAYER
            )
        order = VariableConfigurationService._next_order(role, layer)
        return Variable.objects.create(
            name=name, role=role, order=order, mediator_layer=layer
        )


# ======================================================================
# Contoh format (file template yang diunduh user)
# ======================================================================

VILLAGE_SAMPLE = [
    (1, "Batu", "Oro-oro Ombo"),
    (2, "Batu", "Ngaglik"),
    (9, "Bumiaji", "Bumiaji"),
]

VARIABLE_SAMPLE = [
    ("X1", "Orientasi Pasar", "X1.1", "Kualitas pelayanan wisata",
     "Pengelola wisata memberikan pelayanan yang baik kepada wisatawan.",
     "Skala Likert (1-5)"),
    ("X1", "Orientasi Pasar", "X1.2", "Pemahaman terhadap kebutuhan wisatawan",
     "Pengelola wisata memahami kebutuhan dan keinginan wisatawan.",
     "Skala Likert (1-5)"),
    ("X2", "Fasilitas Pariwisata", "X2.1", "Ketersediaan fasilitas dasar",
     "Fasilitas dasar di desa wisata tersedia dengan memadai.",
     "Skala Likert (1-5)"),
]

DISTRICT_SAMPLE = [
    (1, "Batu"),
    (2, "Bumiaji"),
    (3, "Junrejo"),
]

DISTRICT_HEADERS = ["ID Kecamatan", "Kecamatan"]
VILLAGE_HEADERS = ["ID Desa", "Kecamatan", "Desa/Kelurahan"]
VARIABLE_HEADERS = [
    "Kode Variabel", "Nama Variabel", "Kode Indikator", "Nama Indikator",
    "Pernyataan", "Jenis Jawaban",
]


def _template(headers, sample, notes, title):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sample:
        sheet.append(list(row))
    for column in sheet.columns:
        width = max(len(str(c.value or "")) for c in column)
        sheet.column_dimensions[column[0].column_letter].width = min(width + 2, 70)

    guide = workbook.create_sheet("Petunjuk")
    guide.append([title])
    guide["A1"].font = Font(bold=True)
    for note in notes:
        guide.append([note])
    guide.column_dimensions["A"].width = 110
    return workbook


def district_template():
    return _template(
        DISTRICT_HEADERS,
        DISTRICT_SAMPLE,
        [
            "Data dibaca dari sheet PERTAMA (Sheet1); sheet ini diabaikan.",
            "Baris 1 adalah header. Isi data mulai baris 2.",
            "ID Kecamatan: angka unik per kecamatan (1, 2, 3, ...).",
            "Kecamatan: nama kecamatan, unik.",
            "Impor kecamatan sebelum desa agar ID Kecamatan sesuai file.",
            "Kolom lain di luar A-B diabaikan.",
        ],
        "Petunjuk impor Kecamatan",
    )


def village_template():
    return _template(
        VILLAGE_HEADERS,
        VILLAGE_SAMPLE,
        [
            "Data dibaca dari sheet PERTAMA (Sheet1); sheet ini diabaikan.",
            "Baris 1 adalah header. Isi data mulai baris 2.",
            "ID Desa: angka unik per desa (1, 2, 3, ...).",
            "Kecamatan: nama kecamatan. Jika belum ada, dibuat otomatis (impor Kecamatan dulu bila ID-nya harus sesuai file).",
            "Desa/Kelurahan: nama desa, unik.",
            "Kolom lain di luar A-C diabaikan.",
        ],
        "Petunjuk impor Desa",
    )


def variable_template():
    return _template(
        VARIABLE_HEADERS,
        VARIABLE_SAMPLE,
        [
            "Data dibaca dari sheet PERTAMA (Sheet1); sheet ini diabaikan.",
            "Baris 1 adalah header. Isi data mulai baris 2.",
            "Satu baris = satu pernyataan. Variabel dan indikator yang sama boleh "
            "muncul di beberapa baris.",
            "Kode Variabel: huruf X (predictor), Y (mediator), atau Z (response) "
            "diikuti angka, mis. X1, Y2, Z1.",
            "Kode Indikator: kode variabel + titik + nomor, mis. X1.1. Harus "
            "cocok dengan Kode Variabel di baris yang sama.",
            "Nama Variabel dan Nama Indikator harus unik; maksimal 100 karakter.",
            "Pernyataan menjadi pertanyaan kuesioner.",
            "Jenis Jawaban: " + ", ".join(ANSWER_TYPE_LABELS.values()) + ".",
            "Variabel Y masuk ke Mediator Layer 1; pindahkan lewat Configure "
            "Variables bila perlu.",
            "Kode final dibuat ulang sistem berdasarkan role dan urutan.",
        ],
        "Petunjuk impor Variabel, Indikator, dan Pernyataan",
    )
