"""Pembacaan Excel sumber + Schema Gate (belum menulis ke database)."""

import hashlib
import os
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook

from apps.master.models import Indicator, Questionnaire, Village
from apps.response.services.bulk_import_service import (
    ResponseBulkImportService,
)

from .rules import IDENTITY_COLUMNS, display_name
from .wording import MSG_NO_QUESTIONNAIRE, msg_missing_columns


@dataclass
class SourceRow:
    row_number: int
    values: dict


@dataclass
class SourceData:
    header_row: int
    column_map: dict
    rows: list = field(default_factory=list)


class SchemaError(Exception):
    def __init__(self, missing, message):
        super().__init__(message)
        self.missing = missing


def is_blank(value):
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    if isinstance(value, float):
        return value != value  # NaN
    return False


def as_integer(value):
    """int bila value bilangan bulat numerik (bukan bool/teks), selain itu None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return None


def read_bytes(file):
    """Return (bytes, nama_file) dari path, bytes, atau file-like."""
    if isinstance(file, (bytes, bytearray)):
        return bytes(file), ""
    if isinstance(file, (str, os.PathLike)):
        with open(file, "rb") as handle:
            return handle.read(), os.path.basename(str(file))
    return file.read(), os.path.basename(getattr(file, "name", "") or "")


def sha256_of(data):
    return hashlib.sha256(data).hexdigest()


def read_source(data):
    """
    Baca sheet aktif memakai pencari header yang sama dengan import
    existing. Baris yang seluruh selnya kosong tidak dihitung sebagai
    rekaman. Nilai sel dibiarkan apa adanya (tidak dikonversi).

    Raises BulkImportError bila header "ID Resp"/"Desa" tidak ada.
    """

    wb = load_workbook(filename=BytesIO(data), data_only=True)
    ws = wb.active

    header_row, column_map = ResponseBulkImportService._find_header_columns(ws)

    source = SourceData(header_row=header_row, column_map=column_map)

    for row_number in range(header_row + 1, ws.max_row + 1):
        values = {
            name: ws.cell(row=row_number, column=col).value
            for name, col in column_map.items()
        }
        if all(is_blank(v) for v in values.values()):
            continue
        source.rows.append(SourceRow(row_number, values))

    return source


def expected_indicator_codes():
    """
    Kode indikator wajib = indikator dari kuesioner aktif di master
    (sumber yang sama dengan import existing), huruf kecil, terurut.
    """

    codes = (
        Questionnaire.objects
        .filter(is_active=True)
        .values_list("indicator__code", flat=True)
    )
    return sorted({code.lower() for code in codes})


def indicators_without_questionnaire():
    """Kode indikator aktif yang belum punya kuesioner aktif (urut)."""

    covered = set(
        Questionnaire.objects
        .filter(is_active=True)
        .values_list("indicator_id", flat=True)
    )
    return sorted(
        code
        for pk, code in Indicator.objects.filter(is_active=True)
        .values_list("id", "code")
        if pk not in covered
    )


def registered_village_codes():
    return set(Village.objects.values_list("code", flat=True))


def schema_gate(column_map, indicator_codes):
    """
    Pastikan seluruh kolom wajib ada. Tidak menebak / membuat kolom.
    Raises SchemaError dengan daftar kolom yang hilang.
    """

    if not indicator_codes:
        raise SchemaError([], MSG_NO_QUESTIONNAIRE)

    required = list(IDENTITY_COLUMNS) + list(indicator_codes)
    missing = [name for name in required if name not in column_map]

    if missing:
        names = [display_name(name) for name in missing]
        raise SchemaError(names, msg_missing_columns(names))
