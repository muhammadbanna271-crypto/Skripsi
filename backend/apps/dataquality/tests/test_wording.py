from django.contrib.auth import get_user_model
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.dataquality.pipeline import ValidatedImportService
from apps.dataquality.wording import result_messages
from apps.master.models import Indicator, Questionnaire, Variable

from .helpers import (
    COL_DESA,
    COL_ID,
    FIRST_INDICATOR_COL,
    INDICATOR_CODES,
    DQTestCase,
    make_workbook,
)


def set_cell(row, col, value):
    def mutate(ws):
        ws.cell(row=row, column=col).value = value
    return mutate


TECHNICAL_TERMS = ("null", "NaN", "P1", "P2", "checkpoint", "rule",
                   "Sel wajib", "Rekaman", "dataset")


class PlainLanguageTests(DQTestCase):

    def run_import(self, data):
        return ValidatedImportService.run(data, self.survey)

    def texts(self, execution):
        return [text for _, text in result_messages(execution)]

    def reasons(self, execution):
        return [r.reason for r in execution.rejections.all()]

    # ---- kegagalan struktur ------------------------------------
    def test_no_questionnaire_explains_what_to_do(self):
        Questionnaire.objects.all().delete()

        execution = self.run_import(make_workbook())

        self.assertEqual(execution.status, "failed_schema")
        message = self.texts(execution)[0]
        self.assertIn("Belum ada kuesioner aktif", message)
        self.assertIn("menu Questionnaire", message)
        self.assertIn("Tidak ada data yang disimpan", message)
        self.assertIn(execution.execution_id, message)
        self.assertNotIn("Master", message)

    def test_missing_columns_message_names_columns(self):
        def drop(ws):
            ws.delete_cols(FIRST_INDICATOR_COL + 3)

        execution = self.run_import(make_workbook(drop))

        message = self.texts(execution)[0]
        self.assertIn("1 kolom wajib tidak ditemukan", message)
        self.assertIn(INDICATOR_CODES[3], message)
        self.assertIn("Struktur file tidak sesuai", message)

    def test_many_missing_columns_are_shortened(self):
        def drop_many(ws):
            ws.delete_cols(FIRST_INDICATOR_COL, 30)

        execution = self.run_import(make_workbook(drop_many))

        message = self.texts(execution)[0]
        self.assertIn("30 kolom wajib tidak ditemukan", message)
        self.assertIn("dan 20 lainnya", message)

    def test_unreadable_file_message(self):
        execution = self.run_import(b"ini bukan file excel")

        self.assertEqual(execution.status, "failed_process")
        self.assertIn("File tidak dapat dibuka", self.texts(execution)[0])
        self.assertNotIn("BadZipFile", self.texts(execution)[0])

    def test_header_not_found_message(self):
        from openpyxl import Workbook
        from io import BytesIO

        wb = Workbook()
        wb.active.append(["nama", "umur"])
        buffer = BytesIO()
        wb.save(buffer)

        execution = self.run_import(buffer.getvalue())

        self.assertEqual(execution.status, "failed_schema")
        self.assertIn("Judul kolom pada file tidak dapat dikenali",
                      self.texts(execution)[0])

    # ---- peringatan indikator tanpa kuesioner ------------------
    def test_warns_about_indicator_without_questionnaire(self):
        variable = Variable.objects.get(code="X1")
        Indicator.objects.create(
            code="X1.99", name="Tanpa kuesioner", variable=variable
        )

        execution = self.run_import(make_workbook())

        self.assertEqual(execution.status, "success")
        warning = [t for t in self.texts(execution) if "X1.99" in t]
        self.assertEqual(len(warning), 1)
        self.assertIn("belum memiliki kuesioner aktif", warning[0])

    # ---- alasan penolakan --------------------------------------
    def test_rejection_reasons_are_plain_sentences(self):
        def mutate(ws):
            ws.cell(row=3, column=FIRST_INDICATOR_COL).value = None
            ws.cell(row=4, column=FIRST_INDICATOR_COL).value = "3"
            ws.cell(row=5, column=FIRST_INDICATOR_COL).value = 7
            ws.cell(row=6, column=COL_DESA).value = 999
            ws.cell(row=10, column=COL_ID).value = 2

        execution = self.run_import(make_workbook(mutate))
        by_rule = {r.rule: r.reason for r in execution.rejections.all()}

        self.assertIn("masih kosong", by_rule["required_cell_missing"])
        self.assertIn("angka bulat 1 sampai 5", by_rule["indicator_integer"])
        self.assertIn("tertulis bukan angka bulat: 3",
                      by_rule["indicator_integer"])
        self.assertIn("antara 1 sampai 5", by_rule["indicator_range"])
        self.assertIn("tidak terdaftar di data master desa",
                      by_rule["village_not_registered"])
        self.assertIn("sudah dipakai di baris 3",
                      by_rule["id_resp_duplicate"])

        for reason in self.reasons(execution):
            for term in TECHNICAL_TERMS:
                self.assertNotIn(term, reason)

    # ---- ringkasan hasil ---------------------------------------
    def test_success_summary(self):
        execution = self.run_import(make_workbook())

        messages = result_messages(execution)
        self.assertEqual(messages[0][0], "success")
        self.assertIn("24 baris berhasil disimpan", messages[0][1])
        self.assertEqual(len(messages), 1)

    def test_rejection_summary_groups_problems(self):
        def mutate(ws):
            ws.cell(row=3, column=FIRST_INDICATOR_COL).value = None
            ws.cell(row=4, column=FIRST_INDICATOR_COL).value = None
            ws.cell(row=5, column=FIRST_INDICATOR_COL).value = 9

        execution = self.run_import(make_workbook(mutate))
        texts = self.texts(execution)

        self.assertIn("21 baris berhasil disimpan dan 3 baris tidak "
                      "diproses", texts[0])
        self.assertIn("2 baris: ada isian yang masih kosong", texts[1])
        self.assertIn("1 baris: ada angka di luar rentang 1 sampai 5",
                      texts[1])
        self.assertIn("Baris 3, kolom", texts[1])
        self.assertIn("tidak memperbaiki atau mengisi data secara otomatis",
                      texts[2])

    # ---- tampilan di halaman impor -----------------------------
    def test_import_page_shows_plain_messages(self):
        user = get_user_model().objects.create_user(
            "staff", password="x", is_staff=True
        )
        self.client.force_login(user)

        response = self.client.post(
            reverse("response:response-import"),
            {
                "survey_id": self.survey.pk,
                "excel_file": SimpleUploadedFile(
                    "a.xlsx",
                    make_workbook(set_cell(5, FIRST_INDICATOR_COL, "x")),
                ),
            },
        )

        shown = [str(m) for m in get_messages(response.wsgi_request)]
        self.assertTrue(any("Impor selesai" in m for m in shown))
        self.assertTrue(any("perlu diperbaiki di file Excel" in m
                            for m in shown))
