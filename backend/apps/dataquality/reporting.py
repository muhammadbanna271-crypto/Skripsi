"""
Ekspor satu eksekusi ke file biasa (CSV + report.txt + fingerprints.txt)
supaya hasil pengujian dapat diperiksa tanpa Django/database.

Hanya membaca database. Database tetap penyimpanan internal.
"""

import csv
import os
from collections import defaultdict

from .evaluation import (
    FINGERPRINT_KEYS,
    compare_executions,
    evaluate,
    scenario_name,
)
from .models import ExecutionLog
from .rules import DIMENSIONS

NOT_COMPUTABLE = "tidak dapat dihitung (n_checked=0)"

SUMMARY_FIELDS = [
    "execution_id", "status", "started_at", "finished_at",
    "source_filename", "source_sha256", "rules_version",
    "n_input", "n_accepted", "n_rejected", "n_rejected_p1",
    "n_rejected_p2", "n_output", "n_lost", "delta_n",
    "person_domain_rule", "error_message",
]

REJECTION_FIELDS = [
    "execution_id", "checkpoint", "source_row", "id_resp", "column_name",
    "original_value", "value_type", "dimension", "rule", "reason",
    "db_value", "related_row", "record_action",
]

QUALITY_FIELDS = [
    "execution_id", "checkpoint", "dimension", "n_checked", "n_satisfied",
    "score", "note",
]

TRUTH_FIELDS = [
    "scenario", "seed", "source_row", "id_resp", "column_name",
    "defect_type", "original_value", "injected_value",
]

EVALUATION_FIELDS = [
    "scenario", "unit", "seed", "tp", "fp", "fn", "precision", "recall",
]

LOCATION_FIELDS = ["scenario", "outcome", "location"]


def _write_csv(path, fields, rows):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _fmt(value):
    return "" if value is None else value


def _pct(score):
    return "n/a" if score is None else f"{score:.4f}"


def rejection_rows(execution, checkpoint=None):
    queryset = execution.rejections.all()
    if checkpoint:
        queryset = queryset.filter(checkpoint=checkpoint)
    return [
        {
            "execution_id": execution.execution_id,
            "checkpoint": r.checkpoint,
            "source_row": r.source_row,
            "id_resp": r.respondent_id,
            "column_name": r.column_name,
            "original_value": r.original_value,
            "value_type": r.value_type,
            "dimension": r.dimension,
            "rule": r.rule,
            "reason": r.reason,
            "db_value": r.db_value,
            "related_row": _fmt(r.related_row),
            "record_action": r.record_action,
        }
        for r in queryset.order_by("checkpoint", "source_row", "id")
    ]


def quality_rows(execution):
    rows = []
    for q in execution.quality_scores.order_by("checkpoint", "id"):
        rows.append(
            {
                "execution_id": execution.execution_id,
                "checkpoint": q.checkpoint,
                "dimension": q.dimension,
                "n_checked": q.n_checked,
                "n_satisfied": q.n_satisfied,
                "score": _fmt(q.score),
                "note": NOT_COMPUTABLE if q.score is None else "",
            }
        )
    order = {d: i for i, d in enumerate(DIMENSIONS)}
    rows.sort(key=lambda r: (r["checkpoint"], order[r["dimension"]]))
    return rows


def truth_rows(truth):
    scenario = scenario_name(truth)
    return [
        {
            "scenario": scenario,
            "seed": truth.get("seed"),
            "source_row": r["row"],
            "id_resp": _fmt(r["respondent_id"]),
            "column_name": r["column"],
            "defect_type": r["defect_type"],
            "original_value": _fmt(r["original_value"]),
            "injected_value": _fmt(r["injected_value"]),
        }
        for r in truth.get("records", [])
    ]


def fingerprint_text(execution):
    summary = execution.output_summary
    state = summary.get("database_state")
    lines = [
        "FINGERPRINT EKSEKUSI",
        "====================",
        f"Execution ID : {execution.execution_id}",
        f"Source SHA256: {execution.source_sha256}",
        f"Rules version: {execution.rules_version}",
        f"Person rule  : {summary.get('config', {}).get('person_domain_rule')}",
        "",
        "Fingerprint dihitung dari KEADAAN database di akhir eksekusi",
        "(seluruh tabel, bukan hanya baris eksekusi ini), tanpa timestamp",
        "dan tanpa primary key. Dua fingerprint hanya bermakna jika sumber,",
        "versi aturan, dan konfigurasi sama.",
        "",
    ]

    if not state:
        lines.append(
            "Tidak tersedia: eksekusi tidak mencapai tahap akhir "
            f"(status: {execution.status})."
        )
        return "\n".join(lines) + "\n"

    for hash_key, count_key, label in FINGERPRINT_KEYS:
        lines += [
            f"[{label.upper()}]",
            f"  jumlah record: {state.get(count_key)}",
            f"  sha256       : {state.get(hash_key)}",
            "",
        ]

    if summary.get("scored_villages") is None:
        lines.append(
            "Catatan: agregasi skor tidak dijalankan pada eksekusi ini "
            "(run_scoring=False); fingerprint skor mencerminkan tabel skor "
            "apa adanya."
        )
    return "\n".join(lines) + "\n"


def _quality_block(execution, checkpoint):
    scores = {
        q.dimension: q
        for q in execution.quality_scores.filter(checkpoint=checkpoint)
    }
    if not scores:
        return ["  (tidak dihitung)"]
    lines = []
    for dimension in DIMENSIONS:
        q = scores.get(dimension)
        if q is None:
            lines.append(f"  {dimension.capitalize():<13}: -")
        elif q.score is None:
            lines.append(
                f"  {dimension.capitalize():<13}: {NOT_COMPUTABLE}"
            )
        else:
            lines.append(
                f"  {dimension.capitalize():<13}: {q.score:.4f}  "
                f"({q.n_satisfied}/{q.n_checked})"
            )
    return lines


def _rejection_summary(execution):
    stats = defaultdict(lambda: [0, set()])
    for r in execution.rejections.all():
        entry = stats[(r.checkpoint, r.dimension)]
        entry[0] += 1
        entry[1].add(r.source_row)

    lines = []
    for checkpoint in ("P1", "P2"):
        lines.append(f"  {checkpoint}  (pelanggaran / baris terdampak)")
        for dimension in DIMENSIONS:
            n_viol, rows = stats.get((checkpoint, dimension), (0, set()))
            lines.append(
                f"    {dimension.capitalize():<13}: {n_viol} / {len(rows)}"
            )
    return lines


def _uniqueness_detail(execution):
    lines = []
    for r in execution.rejections.filter(
        checkpoint="P1", dimension="uniqueness"
    ).order_by("source_row", "id")[:50]:
        lines.append(
            f"  baris {r.source_row} (ID Resp {r.respondent_id}) DITOLAK "
            f"[{r.rule}]; kemunculan pertama dipertahankan pada baris "
            f"{r.related_row}"
        )
    return lines or ["  (tidak ada duplikasi terdeteksi)"]


def _discrepancy_detail(execution, limit=50):
    queryset = execution.rejections.filter(checkpoint="P2").order_by(
        "source_row", "id"
    )
    lines = []
    for r in queryset[:limit]:
        lines.append(
            f"  baris {r.source_row} (ID Resp {r.respondent_id}) "
            f"{r.column_name}: sumber={r.original_value!r} "
            f"database={r.db_value!r} [{r.rule}] tindakan="
            f"{r.record_action or '-'}"
        )
    extra = queryset.count() - limit
    if extra > 0:
        lines.append(f"  ... {extra} lainnya, lihat discrepancy_detail.csv")
    return lines or ["  (tidak ada discrepancy)"]


def report_text(execution, truth=None, compare_with=None):
    summary = execution.output_summary
    config = summary.get("config", {})
    failed_schema = execution.status == ExecutionLog.STATUS_FAILED_SCHEMA

    L = []
    add = L.append

    add("EXECUTION")
    add("---------")
    add(f"Execution ID : {execution.execution_id}")
    add(f"Status       : {execution.status} ({execution.get_status_display()})")
    add(f"Mulai        : {execution.started_at}")
    add(f"Selesai      : {_fmt(execution.finished_at)}")
    add(f"File sumber  : {execution.source_filename or '-'}")
    add(f"SHA-256      : {execution.source_sha256 or '-'}")
    add(f"Versi aturan : {execution.rules_version}")
    add(f"Input        : {execution.n_input}")
    add(f"Accepted (P1): {execution.n_accepted}")
    add(f"Rejected     : {execution.n_rejected}")
    add(f"Output       : {execution.n_output}")
    add(f"Lost         : {execution.n_lost}")
    add(f"Delta N      : {_fmt(execution.delta_n) if execution.delta_n is not None else 'n/a'}")

    if execution.error_message:
        add("")
        add("ERROR")
        add("-----")
        add(execution.error_message)

    add("")
    add("CONFIGURATION")
    add("-------------")
    if config.get("person_domain_rule") == "active":
        add(
            "Validasi domain Person : AKTIF, nilai diizinkan = "
            f"{config.get('allowed_persons')}"
        )
    else:
        add(
            "Validasi domain Person : TIDAK AKTIF (DATAQUALITY_ALLOWED_"
            "PERSONS belum diisi; repository tidak punya daftar jabatan"
        )
        add(
            "                         resmi). Person hanya diperiksa "
            "terisi dan bukan boolean -- domain Person BELUM diuji."
        )

    if failed_schema:
        add("")
        add("SCHEMA GATE")
        add("-----------")
        add("GAGAL. Dataset dihentikan; tidak ada P1/P2 dan tidak ada")
        add("quality score (tidak dihitung, bukan 0).")
        missing = summary.get("missing_columns")
        if missing:
            add(f"Kolom wajib hilang ({len(missing)}): {', '.join(missing)}")
        add(f"Baris data terbaca: {execution.n_input}")
    else:
        add("")
        add("P1 QUALITY  (Q = n_satisfied / n_checked)")
        add("----------")
        L.extend(_quality_block(execution, "P1"))
        add("")
        add("P2 QUALITY")
        add("----------")
        L.extend(_quality_block(execution, "P2"))

        add("")
        add("REJECTION SUMMARY")
        add("-----------------")
        L.extend(_rejection_summary(execution))

        add("")
        add("UNIQUENESS DETAIL (P1)")
        add("----------------------")
        L.extend(_uniqueness_detail(execution))

        add("")
        add("P2 DISCREPANCY (sumber vs database)")
        add("-----------------------------------")
        L.extend(_discrepancy_detail(execution))
        removed = summary.get("p2_removed_niks") or []
        if removed:
            add(
                f"  Rekaman dihapus dari tabel Respondent/Response agar tidak "
                f"ikut skoring ({len(removed)}): {', '.join(removed[:20])}"
                + (" ..." if len(removed) > 20 else "")
            )

        add("")
        add("RECONCILIATION")
        add("--------------")
        add(f"N_input             : {execution.n_input}")
        add(f"N_accepted (lolos P1): {execution.n_accepted}")
        add(f"N_rejected_P1       : {execution.n_rejected_p1}  (ditolak validasi sumber)")
        add(f"N_rejected_P2       : {execution.n_rejected_p2}  (tersimpan tetapi berbeda dari sumber; dikeluarkan dari tabel skoring)")
        add(f"N_rejected          : {execution.n_rejected}  (= P1 + P2)")
        add(f"N_lost              : {execution.n_lost}  (lolos P1, tidak ditemukan di database; BUKAN penolakan)")
        add(f"N_output            : {execution.n_output}  (tersimpan dan sama dengan sumber)")
        add(f"Delta N             : {_fmt(execution.delta_n)}")
        add("Rumus: Delta N = N_input - (N_output + N_rejected)")
        add("       N_output = N_accepted - N_rejected_P2 - N_lost")
        add("       => Delta N = N_lost  (0 berarti tidak ada record hilang)")
        check = execution.n_input - (execution.n_output + execution.n_rejected)
        add(
            "Verifikasi ulang   : "
            + ("cocok" if check == execution.delta_n else "TIDAK COCOK")
        )

    add("")
    add("FINGERPRINT")
    add("-----------")
    state = summary.get("database_state")
    if state:
        for hash_key, count_key, label in FINGERPRINT_KEYS:
            add(
                f"{label:<11}: {state.get(hash_key)}  "
                f"({state.get(count_key)} record)"
            )
    else:
        add("Tidak tersedia (eksekusi tidak mencapai tahap akhir).")

    if truth is not None:
        add("")
        add("DETECTION vs GROUND TRUTH")
        add("-------------------------")
        result = evaluate(execution, truth)
        add(f"Skenario : {result['scenario']}  (seed {result['seed']})")
        add(f"Satuan lokasi cacat: {result['unit']}")
        add(f"TP={result['tp']}  FP={result['fp']}  FN={result['fn']}")
        add(f"Precision: {_pct(result['precision'])}")
        if result["unit"] == "clean":
            add("Recall   : n/a (dataset bersih -- fokus pada false positive)")
        else:
            add(f"Recall   : {_pct(result['recall'])}")
        for name in ("fp", "fn"):
            locations = result[f"{name}_locations"]
            if locations:
                add(f"Lokasi {name.upper()}: {locations[:20]}")
        add(
            "Catatan: satu cacat bisa melanggar >1 rule; evaluasi memakai "
            "satuan lokasi di atas sehingga tidak dihitung ganda."
        )

    if compare_with is not None:
        add("")
        add(f"COMPARISON dengan {compare_with.execution_id}")
        add("-" * 60)
        L.extend(comparison_lines(execution, compare_with))

    return "\n".join(L) + "\n"


def comparison_lines(a, b):
    result = compare_executions(a, b)
    context = result["context"]
    lines = [
        f"A: {a.execution_id}",
        f"B: {b.execution_id}",
        "Konteks:",
        f"  sumber (SHA-256) sama : {context['same_source_sha256']}",
        f"  versi aturan sama     : {context['same_rules_version']}",
        f"  konfigurasi sama      : {context['same_config']}",
        f"  jumlah (in/acc/rej/lost/out) sama: {context['same_counts']}",
    ]
    if not result["comparable"]:
        lines.append(
            "TIDAK DAPAT DIBANDINGKAN: salah satu eksekusi tidak berstatus "
            "success."
        )
        return lines
    for c in result["components"]:
        lines.append(
            f"  {c['label']:<10}: {'SAMA' if c['same'] else 'BERBEDA'}  "
            f"(A: {c['count_a']} rec {str(c['hash_a'])[:12]}..., "
            f"B: {c['count_b']} rec {str(c['hash_b'])[:12]}...)"
        )
    verdict = "SAMA" if result["all_same"] else "BERBEDA"
    lines.append(f"HASIL: {verdict}")
    if not context["same_source_sha256"]:
        lines.append(
            "Peringatan: sumber berbeda; perbedaan fingerprint wajar dan "
            "bukan bukti ketidakdeterministikan."
        )
    lines.append(
        "Tafsir: jika database direset di antara A dan B -> uji "
        "determinisme; jika tidak direset (upload ulang) -> uji idempotensi."
    )
    return lines


def build_workbook(execution):
    """Satu workbook berisi ringkasan, skor kualitas, dan daftar penolakan."""

    from openpyxl import Workbook
    from openpyxl.styles import Font

    summary = execution.output_summary
    workbook = Workbook()

    def fill(sheet, fields, rows):
        sheet.append(fields)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([row.get(field, "") for field in fields])
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column[:200])
            sheet.column_dimensions[column[0].column_letter].width = min(
                max(width + 2, 10), 60
            )

    ringkasan = workbook.active
    ringkasan.title = "Ringkasan"
    fill(
        ringkasan,
        SUMMARY_FIELDS,
        [
            {
                "execution_id": execution.execution_id,
                "status": execution.status,
                "started_at": str(execution.started_at),
                "finished_at": str(_fmt(execution.finished_at)),
                "source_filename": execution.source_filename,
                "source_sha256": execution.source_sha256,
                "rules_version": execution.rules_version,
                "n_input": execution.n_input,
                "n_accepted": execution.n_accepted,
                "n_rejected": execution.n_rejected,
                "n_rejected_p1": execution.n_rejected_p1,
                "n_rejected_p2": execution.n_rejected_p2,
                "n_output": execution.n_output,
                "n_lost": execution.n_lost,
                "delta_n": _fmt(execution.delta_n),
                "person_domain_rule": summary.get("config", {}).get(
                    "person_domain_rule", ""
                ),
                "error_message": execution.error_message,
            }
        ],
    )
    fill(
        workbook.create_sheet("Skor Kualitas"),
        QUALITY_FIELDS,
        quality_rows(execution),
    )
    fill(
        workbook.create_sheet("Ditolak (P1+P2)"),
        REJECTION_FIELDS,
        rejection_rows(execution),
    )
    fill(
        workbook.create_sheet("Discrepancy (P2)"),
        REJECTION_FIELDS,
        rejection_rows(execution, "P2"),
    )
    return workbook


def export_execution(execution, output_dir, *, truth=None, compare_with=None):
    """Tulis seluruh artefak ke output_dir; return daftar path."""

    os.makedirs(output_dir, exist_ok=True)

    if truth is None:
        truth = execution.output_summary.get("ground_truth")

    summary = execution.output_summary
    paths = {}

    def path(name):
        paths[name] = os.path.join(output_dir, name)
        return paths[name]

    _write_csv(
        path("execution_summary.csv"),
        SUMMARY_FIELDS,
        [
            {
                "execution_id": execution.execution_id,
                "status": execution.status,
                "started_at": execution.started_at,
                "finished_at": _fmt(execution.finished_at),
                "source_filename": execution.source_filename,
                "source_sha256": execution.source_sha256,
                "rules_version": execution.rules_version,
                "n_input": execution.n_input,
                "n_accepted": execution.n_accepted,
                "n_rejected": execution.n_rejected,
                "n_rejected_p1": execution.n_rejected_p1,
                "n_rejected_p2": execution.n_rejected_p2,
                "n_output": execution.n_output,
                "n_lost": execution.n_lost,
                "delta_n": _fmt(execution.delta_n),
                "person_domain_rule": summary.get("config", {}).get(
                    "person_domain_rule", ""
                ),
                "error_message": execution.error_message,
            }
        ],
    )

    _write_csv(
        path("rejection_detail.csv"), REJECTION_FIELDS,
        rejection_rows(execution),
    )
    _write_csv(
        path("discrepancy_detail.csv"), REJECTION_FIELDS,
        rejection_rows(execution, "P2"),
    )
    _write_csv(path("quality_score.csv"), QUALITY_FIELDS, quality_rows(execution))

    with open(path("fingerprints.txt"), "w", encoding="utf-8") as handle:
        handle.write(fingerprint_text(execution))

    if truth is not None:
        _write_csv(path("ground_truth.csv"), TRUTH_FIELDS, truth_rows(truth))

        result = evaluate(execution, truth)
        _write_csv(
            path("evaluation.csv"),
            EVALUATION_FIELDS,
            [{k: _fmt(result[k]) for k in EVALUATION_FIELDS}],
        )
        _write_csv(
            path("evaluation_locations.csv"),
            LOCATION_FIELDS,
            [
                {
                    "scenario": result["scenario"],
                    "outcome": outcome.upper(),
                    "location": location,
                }
                for outcome in ("tp", "fp", "fn")
                for location in result[f"{outcome}_locations"]
            ],
        )

    if compare_with is not None:
        with open(path("comparison.txt"), "w", encoding="utf-8") as handle:
            handle.write(
                "\n".join(comparison_lines(execution, compare_with)) + "\n"
            )

    with open(path("report.txt"), "w", encoding="utf-8") as handle:
        handle.write(report_text(execution, truth, compare_with))

    return paths
